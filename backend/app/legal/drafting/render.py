"""A draft as Word and PDF (S3.1, DOC-01; decision D-045).

- **Word** is written into the firm's own file when the template version carries one: its styles,
  page set-up, header and margins are kept and its body is replaced. Otherwise a plain document is
  used. Headings use the "Heading 1–3" styles; a gap ([[NEEDS INPUT: …]]) is highlighted.
- **The footer** says what the document is: the matter, the template version, and either the
  approval (who, when, the content hash it bound) or "DRAFT — NOT APPROVED".
- **Metadata is cleared** (author, last editor, comments), as for exports (S1.10).
- **PDF** (reportlab, A4) is the same content, with "DRAFT — NOT APPROVED" across each page until approved.
"""
from __future__ import annotations

import io
import re
from typing import Optional
from xml.sax.saxutils import escape

GAP = re.compile(r"(\[\[NEEDS INPUT:[^\]]*\]\])")


def footer_text(content: dict, approval: Optional[dict], version: Optional[int] = None) -> str:
    t = content["template"]
    head = f"{content['matter']['number']} · {t['name']} v{t['version']}"
    if approval:
        return f"{head} · Approved by {approval['by']} on {approval['at'][:10]} · content {approval['content_sha256'][:12]}…"
    return f"{head} · DRAFT — NOT APPROVED" + (f" (version {version})" if version else "")


def _clear_body(doc) -> None:
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def _runs(par, text: str) -> None:
    from docx.enum.text import WD_COLOR_INDEX
    lines = text.split("\n")
    for i, line in enumerate(lines):
        for piece in GAP.split(line):
            if not piece:
                continue
            run = par.add_run(piece)
            if GAP.fullmatch(piece):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                run.bold = True
        if i < len(lines) - 1:
            par.add_run().add_break()


def to_docx(content: dict, approval: Optional[dict], styles: Optional[bytes] = None, version: Optional[int] = None) -> bytes:
    import docx
    doc = docx.Document(io.BytesIO(styles)) if styles else docx.Document()
    _clear_body(doc)
    names = {s.name for s in doc.styles}
    for b in content["blocks"]:
        if b["type"] == "heading":
            style = f"Heading {b['level']}"
            if style in names:
                _runs(doc.add_paragraph(style=style), b["text"])
            else:
                p = doc.add_paragraph()
                _runs(p, b["text"])
                for r in p.runs:
                    r.bold = True
        else:
            _runs(doc.add_paragraph(), b["text"])
    text = footer_text(content, approval, version)
    for section in doc.sections:
        f = section.footer
        f.is_linked_to_previous = False
        p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
        p.text = text
    cp = doc.core_properties
    cp.author, cp.last_modified_by, cp.comments, cp.keywords, cp.subject = "", "", "", "", ""
    cp.title = content["title"][:200]
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _markup(text: str) -> str:
    parts = []
    for piece in GAP.split(text):
        if GAP.fullmatch(piece or ""):
            parts.append(f'<font backColor="#FFF3A3"><b>{escape(piece)}</b></font>')
        else:
            parts.append(escape(piece))
    return "".join(parts).replace("\n", "<br/>")


def to_pdf(content: dict, approval: Optional[dict], version: Optional[int] = None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontName, body.fontSize, body.leading = "Times-Roman", 11, 15
    heads = {1: styles["Heading1"], 2: styles["Heading2"], 3: styles["Heading3"]}
    story = []
    for b in content["blocks"]:
        if b["type"] == "heading":
            story.append(Paragraph(_markup(b["text"]), heads.get(b["level"], styles["Heading3"])))
        else:
            story += [Paragraph(_markup(b["text"]), body), Spacer(1, 4)]
    foot = footer_text(content, approval, version)

    def page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, 12 * mm, foot)
        canvas.drawRightString(190 * mm, 12 * mm, str(doc.page))
        if not approval:
            canvas.setFont("Helvetica-Bold", 40)
            canvas.setFillGray(0.85)
            canvas.translate(105 * mm, 150 * mm)
            canvas.rotate(45)
            canvas.drawCentredString(0, 0, "DRAFT — NOT APPROVED")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm, bottomMargin=20 * mm,
                      title=content["title"][:200], author="").build(story, onFirstPage=page, onLaterPages=page)
    return buf.getvalue()
