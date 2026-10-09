"""Rendering an approved export (S1.10).

PDF from a PDF: every planned string is found on each page and removed with redaction
annotations that are *applied* (the text under the box is deleted, not covered); annotations,
form fields, attachments, scripts, metadata and XMP are stripped; the watermark goes on every page.

Everything else is rebuilt from the readable content (source.py): a new Word file or a new PDF
with the planned strings replaced by a black "[REDACTED]" mark. Nothing of the original file's
structure — comments, revisions, hidden runs, custom XML, properties — can carry over.

`render(job)` takes {"format", "source", "items" (strings to remove), "watermark"} and returns bytes;
verify.py checks the result before anything is released.
"""
from __future__ import annotations

import html
import io
import re

REDACTED = "[REDACTED]"


def _pattern(items: list[str]):
    items = sorted({i for i in items if i}, key=len, reverse=True)
    return re.compile("|".join(re.escape(i) for i in items), re.I) if items else None


def redact_runs(runs: list[dict], pat) -> list[dict]:
    """Split runs at matches; each match becomes one {"text": "[REDACTED]", "redacted": True} run,
    even when it spans several runs."""
    if pat is None:
        return runs
    full = "".join(r["text"] for r in runs)
    spans = [m.span() for m in pat.finditer(full) if m.end() > m.start()]
    if not spans:
        return runs
    out, pos, emitted = [], 0, set()
    for r in runs:
        start, end = pos, pos + len(r["text"])
        pos, cur = end, start
        for si, (s, e) in enumerate(spans):
            if e <= cur or s >= end:
                continue
            if s > cur:
                out.append({**r, "text": full[cur:s]})
            if si not in emitted:
                out.append({"text": REDACTED, "redacted": True, "b": False, "i": False, "u": False})
                emitted.add(si)
            cur = min(e, end)
        if cur < end:
            out.append({**r, "text": full[cur:end]})
    return out


def _redact_blocks(blocks: list[dict], pat) -> list[dict]:
    out = []
    for b in blocks:
        if b["type"] == "p":
            out.append({**b, "runs": redact_runs(b["runs"], pat)})
        else:
            out.append({"type": "table", "rows": [[[{**p, "runs": redact_runs(p["runs"], pat)} for p in cell] for cell in row]
                                                  for row in b["rows"]]})
    return out


# ── Word ─────────────────────────────────────────────────────────────────────────
def _style(style) -> str:
    m = re.match(r"(?i)heading\s*([1-9])$", style or "")
    if m:
        return f"Heading {m.group(1)}"
    return "Title" if (style or "").lower() == "title" else "Normal"


def _docx_runs(par, runs) -> None:
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import RGBColor
    for r in runs:
        run = par.add_run(r["text"])
        if r.get("redacted"):
            run.font.highlight_color = WD_COLOR_INDEX.BLACK
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        else:
            run.bold, run.italic, run.underline = r.get("b") or None, r.get("i") or None, r.get("u") or None


def to_docx(blocks: list[dict], watermark: str) -> bytes:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc = Document()
    # python-docx's blank template carries custom XML (bibliography sources) and a thumbnail; a
    # clean export carries neither. Parts no relationship reaches are not written.
    for rels, kinds in ((doc.part.rels, (RT.CUSTOM_XML,)), (doc.part.package.rels, (RT.THUMBNAIL,))):
        for rid in [rid for rid, rel in rels.items() if rel.reltype in kinds]:
            del rels[rid]
    cp = doc.core_properties
    cp.author, cp.last_modified_by, cp.title, cp.subject, cp.keywords, cp.comments, cp.category = ("EchoMind Law", "", "", "", "", "", "")
    doc.sections[0].header.paragraphs[0].text = watermark
    doc.sections[0].footer.paragraphs[0].text = watermark
    for b in blocks:
        if b["type"] == "p":
            _docx_runs(doc.add_paragraph(style=_style(b.get("style"))), b["runs"])
        else:
            cols = max((len(r) for r in b["rows"]), default=0)
            if not cols:
                continue
            table = doc.add_table(rows=len(b["rows"]), cols=cols)
            for ri, row in enumerate(b["rows"]):
                for ci, cell in enumerate(row):
                    target = table.cell(ri, ci)
                    target.paragraphs[0].text = ""
                    for pi, p in enumerate(cell):
                        par = target.paragraphs[0] if pi == 0 else target.add_paragraph()
                        _docx_runs(par, p["runs"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────────
_WM_STYLE = "font-family:sans-serif;font-size:7px;color:#555;text-align:center"


def _watermark_pages(doc, watermark: str) -> None:
    import fitz
    for page in doc:
        r = page.rect
        page.insert_htmlbox(fitz.Rect(36, r.height - 26, r.width - 36, r.height - 10), f'<div style="{_WM_STYLE}">{html.escape(watermark)}</div>')


def _scrub(doc) -> None:
    for page in doc:
        for a in list(page.annots() or []):
            page.delete_annot(a)
        for w in list(page.widgets() or []):
            page.delete_widget(w)
    doc.scrub(attached_files=True, clean_pages=True, embedded_files=True, hidden_text=True, javascript=True, metadata=True,
              redactions=True, remove_links=True, reset_fields=True, reset_responses=True, thumbnails=True, xml_metadata=True)
    doc.set_metadata({})
    doc.del_xml_metadata()


def pdf_from_pdf(data: bytes, items: list[str], watermark: str) -> bytes:
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    targets = sorted({i for i in items if i}, key=len, reverse=True)
    for page in doc:
        hits = 0
        for t in targets:
            for rect in page.search_for(t):
                page.add_redact_annot(rect, fill=(0, 0, 0))
                hits += 1
        if hits:
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
    _scrub(doc)
    _watermark_pages(doc, watermark)
    return doc.tobytes(garbage=4, deflate=True, clean=True)


def _html(blocks: list[dict]) -> str:
    def runs(rs):
        out = []
        for r in rs:
            t = html.escape(r["text"]).replace("\t", "&emsp;").replace("\n", "<br/>")
            if r.get("redacted"):
                out.append(t)                         # boxed in black after layout (_box_redactions)
                continue
            for flag, tag in (("u", "u"), ("i", "i"), ("b", "b")):
                if r.get(flag):
                    t = f"<{tag}>{t}</{tag}>"
            out.append(t)
        return "".join(out) or "&nbsp;"
    parts = []
    for b in blocks:
        if b["type"] == "p":
            st = _style(b.get("style"))
            tag = {"Heading 1": "h1", "Heading 2": "h2", "Heading 3": "h3", "Title": "h1"}.get(st, "h3" if st.startswith("Heading") else "p")
            parts.append(f"<{tag}>{runs(b['runs'])}</{tag}>")
        else:
            rows = "".join("<tr>" + "".join(f"<td>{'<br/>'.join(runs(p['runs']) for p in cell)}</td>" for cell in row) + "</tr>"
                           for row in b["rows"])
            parts.append(f'<table border="1" cellpadding="4" style="border-collapse:collapse">{rows}</table>')
    css = "body{font-family:serif;font-size:11pt;line-height:1.4} h1{font-size:16pt} h2{font-size:13pt} h3{font-size:11.5pt} p{margin:0 0 6pt}"
    return f"<style>{css}</style>" + "".join(parts)


def pdf_from_blocks(blocks: list[dict], watermark: str) -> bytes:
    import fitz
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    story = fitz.Story(_html(blocks))
    page, area = fitz.paper_rect("a4"), fitz.paper_rect("a4") + (54, 54, -54, -54)
    more = 1
    while more:
        dev = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(dev)
        writer.end_page()
    writer.close()
    doc = fitz.open(stream=buf.getvalue(), filetype="pdf")
    _box_redactions(doc)
    _scrub(doc)
    _watermark_pages(doc, watermark)
    return doc.tobytes(garbage=4, deflate=True, clean=True)


def _box_redactions(doc) -> None:
    """MuPDF's HTML layout ignores inline backgrounds, so each [REDACTED] mark gets a black box
    drawn over it once the pages are laid out (the mark stays as the page's text)."""
    for page in doc:
        for rect in page.search_for(REDACTED):
            page.draw_rect(rect + (-1, -1, 1, 1), color=(0, 0, 0), fill=(0, 0, 0), overlay=True)


def render(job: dict) -> bytes:
    src, fmt, items, wm = job["source"], job["format"], job["items"], job["watermark"]
    if fmt == "pdf" and src.kind == "pdf":
        return pdf_from_pdf(src.pdf, items, wm)
    blocks = _redact_blocks(src.blocks, _pattern(items))
    return to_docx(blocks, wm) if fmt == "docx" else pdf_from_blocks(blocks, wm)
